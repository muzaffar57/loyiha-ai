from __future__ import annotations

import io

import pytest

pytest.importorskip("ezdxf")
pytest.importorskip("openpyxl")
pytest.importorskip("reportlab")

from openpyxl import load_workbook

from src.formulas import calculate_mass
from src.geometry import analyze_dxf
from src.reports import ReportRow, build_excel, build_pdf
from src.samples import closed_rectangle, dxf_bytes, open_polyline_250


def _row_from_dxf(builder, name: str, d_mm: float, n: int) -> ReportRow:
    part = analyze_dxf(dxf_bytes(builder))
    mass = calculate_mass(part.p_faol_mm, d_mm, n)
    return ReportRow(
        name=name,
        contour_type=part.contour_type,
        p_jami_mm=part.p_jami_mm,
        l_pastki_mm=part.l_pastki_mm,
        p_faol_mm=part.p_faol_mm,
        d_mm=mass.d_mm,
        n=mass.n,
        m_1_kg=mass.m_1_kg,
        waste_kg=mass.waste_kg,
        m_jami_kg=mass.m_jami_kg,
        formula=mass.formula,
        unit_note=part.unit_note,
        warnings=tuple(part.warnings),
        m_without_waste_kg=mass.m_without_waste_kg,
    )


def test_excel_has_part_row_and_totals():
    rectangle = _row_from_dxf(closed_rectangle, "ramka", 3.0, 1)
    opened = _row_from_dxf(open_polyline_250, "ochiq", 3.0, 2)
    blob = build_excel([rectangle, opened])
    wb = load_workbook(io.BytesIO(blob))
    ws = wb.active
    headers = [ws.cell(1, col).value for col in range(1, 11)]
    assert headers[0] == "Detal nomi"
    assert "P_faol" in headers[4]
    assert "M_jami" in headers[9]
    assert ws.cell(2, 1).value == "ramka"
    assert ws.cell(2, 2).value == "yopiq"
    assert ws.cell(2, 3).value == pytest.approx(300)
    assert ws.cell(2, 4).value == pytest.approx(100)
    assert ws.cell(2, 5).value == pytest.approx(200)
    assert ws.cell(2, 8).value == pytest.approx(2.34)
    assert ws.cell(2, 10).value == pytest.approx(2.457)
    assert ws.cell(4, 1).value == "JAMI"
    assert ws.cell(4, 10).value == pytest.approx(rectangle.m_jami_kg + opened.m_jami_kg)


def test_pdf_is_a_real_pdf_with_the_formula():
    row = _row_from_dxf(closed_rectangle, "ramka", 3.0, 1)
    blob = build_pdf([row])
    assert blob.startswith(b"%PDF")
    assert b"2.457" in blob
    assert b"P_faol" in blob
    assert b"M_jami" in blob
    assert b"ramka" in blob
