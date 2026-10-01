"""Excel va PDF hisobotlar. Tashqi xizmat yo'q."""

from __future__ import annotations

import io
from dataclasses import dataclass
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

HEADERS: tuple[str, ...] = (
    "Detal nomi",
    "Kontur turi",
    "P_jami (mm)",
    "L_pastki (mm)",
    "P_faol (mm)",
    "d (mm)",
    "N",
    "M_1 (kg)",
    "Chiqindi 5% (kg)",
    "M_jami (kg)",
)


@dataclass(frozen=True)
class ReportRow:
    name: str
    contour_type: str
    p_jami_mm: float
    l_pastki_mm: float
    p_faol_mm: float
    d_mm: float
    n: int
    m_1_kg: float
    waste_kg: float
    m_jami_kg: float
    formula: str
    unit_note: str = ""
    warnings: tuple[str, ...] = ()
    m_without_waste_kg: float = 0.0


def build_excel(rows: list[ReportRow]) -> bytes:
    """Har bir detal — bitta qator, oxirida JAMI."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Loy sarfi"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E3D")
    thin = Border(
        left=Side(style="thin", color="D0D7D3"),
        right=Side(style="thin", color="D0D7D3"),
        top=Side(style="thin", color="D0D7D3"),
        bottom=Side(style="thin", color="D0D7D3"),
    )
    for col, title in enumerate(HEADERS, start=1):
        cell = ws.cell(1, col, title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin

    for row_index, row in enumerate(rows, start=2):
        values = _row_values(row)
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row_index, col, value)
            cell.border = thin
            if col in (3, 4, 5, 8, 9, 10):
                cell.number_format = "0.000"
            elif col == 6:
                cell.number_format = "0.0"

    total_index = len(rows) + 2
    totals = _totals(rows)
    total_fill = PatternFill("solid", fgColor="E7F2EC")
    bold = Font(bold=True)
    for col, value in enumerate(totals, start=1):
        cell = ws.cell(total_index, col, value)
        cell.font = bold
        cell.fill = total_fill
        cell.border = thin
        if col in (3, 4, 5, 9, 10) and isinstance(value, float):
            cell.number_format = "0.000"

    note_row = total_index + 2
    ws.cell(
        note_row,
        1,
        "Formula: M_1 = (P_faol * d * 2 / 1000) * 1.95;  "
        "M_jami = M_1 * N * 1.05.  "
        "2 — detal uzunligi (m), 1.95 — zichlik (kg/l), 1.05 — 5% chiqindi.  "
        "P_faol: yopiq konturda pastki gorizontal qirra ayiriladi, ochiq konturda ayirilmaydi.",
    )
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=len(HEADERS))
    ws.cell(note_row, 1).alignment = Alignment(wrap_text=True)
    ws.row_dimensions[note_row].height = 48
    ws.row_dimensions[1].height = 30
    ws.auto_filter.ref = f"A1:J{max(total_index, 1)}"
    ws.freeze_panes = "A2"
    widths = (28, 22, 16, 16, 16, 12, 8, 14, 18, 14)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(rows: list[ReportRow]) -> bytes:
    """O'qiladigan PDF. Matn siqilmaydi, shuning uchun formulani ochib ko'rish mumkin."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title="Loy sarfi hisobi",
        pageCompression=0,
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "LoyTitle",
        parent=styles["Title"],
        fontName="Times-Bold",
        fontSize=16,
        textColor=colors.HexColor("#1F4E3D"),
        spaceAfter=6,
    )
    body = ParagraphStyle(
        "LoyBody",
        parent=styles["Normal"],
        fontName="Times-Roman",
        fontSize=10,
        leading=13,
        spaceAfter=4,
    )
    heading = ParagraphStyle(
        "LoyHead",
        parent=styles["Heading2"],
        fontName="Times-Bold",
        fontSize=12,
        textColor=colors.HexColor("#1F4E3D"),
        spaceBefore=8,
        spaceAfter=4,
    )
    small = ParagraphStyle(
        "LoySmall",
        parent=body,
        fontSize=9,
        textColor=colors.HexColor("#333333"),
    )
    story = [
        Paragraph("Loy sarfi hisobi", title),
        Paragraph(
            "Penoplast profiliga qoplama faqat uch yuzga surtiladi: tepa va ikki yon. "
            "Devorga yopishadigan pastki tekis yuz qoplanmaydi.",
            body,
        ),
        Paragraph(
            "M_1 = (P_faol * d * 2 / 1000) * 1.95 kg (bitta 2 m detal). "
            "M_jami = M_1 * N * 1.05 (5% texnik chiqindi).",
            body,
        ),
        Spacer(1, 4 * mm),
    ]
    for row in rows:
        story.append(Paragraph(escape(row.name), heading))
        if row.unit_note:
            story.append(Paragraph(escape(row.unit_note), small))
        for warning in row.warnings:
            story.append(Paragraph("Ogohlantirish: " + escape(warning), small))
        story.append(_part_table(row))
        story.append(Spacer(1, 2 * mm))
        story.append(Paragraph("Formula (sonlar qo'yilgan):", small))
        story.append(
            Preformatted(row.formula, ParagraphStyle(
                "LoyFormula",
                fontName="Courier",
                fontSize=8,
                leading=11,
            ))
        )
    story.append(Paragraph("Jami partiya", heading))
    story.append(_totals_table(rows))
    base = sum(row.m_without_waste_kg for row in rows)
    jami = sum(row.m_jami_kg for row in rows)
    story.append(Spacer(1, 2 * mm))
    story.append(
        Preformatted(
            "Partiya:\n"
            f"M_1 * N yig'indisi = {base:.3f} kg\n"
            f"+5% chiqindi: {base:.3f} * 1.05 = {jami:.3f} kg",
            ParagraphStyle("LoyTotalFormula", fontName="Courier", fontSize=8, leading=11),
        )
    )
    doc.build(story)
    return buf.getvalue()


def _part_table(row: ReportRow) -> Table:
    data = [
        ["Ko'rsatkich", "Qiymat"],
        ["Kontur turi", row.contour_type],
        ["Umumiy perimetr P_jami, mm", f"{row.p_jami_mm:.3f}"],
        ["Pastki chiziq L_pastki, mm", f"{row.l_pastki_mm:.3f}"],
        ["Faol perimetr P_faol, mm", f"{row.p_faol_mm:.3f}"],
        ["Qalinlik d, mm", f"{row.d_mm:.1f}"],
        ["Dona soni N", str(row.n)],
        ["1 dona uchun loy M_1, kg", f"{row.m_1_kg:.3f}"],
        ["M_1 * N, kg", f"{row.m_without_waste_kg:.3f}"],
        ["Chiqindi 5%, kg", f"{row.waste_kg:.3f}"],
        ["Jami loy M_jami, kg", f"{row.m_jami_kg:.3f}"],
    ]
    table = Table(data, colWidths=(78 * mm, 90 * mm))
    table.setStyle(_table_style())
    return table


def _totals_table(rows: list[ReportRow]) -> Table:
    base = sum(row.m_without_waste_kg for row in rows)
    waste = sum(row.waste_kg for row in rows)
    jami = sum(row.m_jami_kg for row in rows)
    data = [
        ["Ko'rsatkich", "Qiymat"],
        ["Detallar soni", str(len(rows))],
        ["P_jami yig'indisi, mm", f"{sum(row.p_jami_mm for row in rows):.3f}"],
        ["L_pastki yig'indisi, mm", f"{sum(row.l_pastki_mm for row in rows):.3f}"],
        ["P_faol yig'indisi, mm", f"{sum(row.p_faol_mm for row in rows):.3f}"],
        ["Jami dona N", str(sum(row.n for row in rows))],
        ["M_1 * N yig'indisi, kg", f"{base:.3f}"],
        ["Chiqindi 5%, kg", f"{waste:.3f}"],
        ["Jami loy M_jami, kg", f"{jami:.3f}"],
    ]
    table = Table(data, colWidths=(78 * mm, 90 * mm))
    table.setStyle(_table_style())
    return table


def _table_style() -> TableStyle:
    return TableStyle(
        [
            ("FONTNAME", (0, 0), (-1, 0), "Times-Bold"),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E3D")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 1), (-1, -1), "Times-Roman"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#F7FBF8")),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#C5D5CC")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
    )


def _row_values(row: ReportRow) -> list[object]:
    return [
        row.name,
        row.contour_type,
        row.p_jami_mm,
        row.l_pastki_mm,
        row.p_faol_mm,
        row.d_mm,
        row.n,
        row.m_1_kg,
        row.waste_kg,
        row.m_jami_kg,
    ]


def _totals(rows: list[ReportRow]) -> list[object]:
    thicknesses = {row.d_mm for row in rows}
    thickness: object
    if len(thicknesses) == 1:
        thickness = next(iter(thicknesses))
    elif not rows:
        thickness = ""
    else:
        thickness = "har xil"
    return [
        "JAMI",
        f"{len(rows)} ta detal",
        sum(row.p_jami_mm for row in rows),
        sum(row.l_pastki_mm for row in rows),
        sum(row.p_faol_mm for row in rows),
        thickness,
        sum(row.n for row in rows),
        "",
        sum(row.waste_kg for row in rows),
        sum(row.m_jami_kg for row in rows),
    ]
